from __future__ import annotations

import csv
import hashlib
from pathlib import Path
from typing import Any, Sequence

from .excel_report import SHEET_CELL_TYPES, SHEET_GEP, SHEET_TISSUES


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _workbook_api():
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Validation requires openpyxl; install requirements.txt") from exc
    return load_workbook


def _forbidden_headers(headers: Sequence[Any]) -> list[str]:
    forbidden = ("ihc", "protein", "candidate", "overall score", "specificity score")
    return [
        str(value)
        for value in headers
        if any(term in str(value or "").lower() for term in forbidden)
    ]


def _sheet_shape_and_headers(worksheet) -> tuple[int, int, list[Any]]:
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = list(next(iterator))
    except StopIteration:
        return 0, 0, []
    rows = sum(1 for _ in iterator)
    return rows, len(headers), headers


def validate_step1(path: Path, expected_rows: int) -> dict[str, Any]:
    errors: list[str] = []
    details: dict[str, Any] = {"path": str(path), "exists": path.exists()}
    if not path.exists():
        return {"status": "FAIL", "errors": [f"Missing workbook: {path}"], "workbook": details}
    workbook = _workbook_api()(path, read_only=True, data_only=True)
    expected_sheets = [SHEET_CELL_TYPES, SHEET_TISSUES]
    details["sheets"] = workbook.sheetnames
    if workbook.sheetnames != expected_sheets:
        errors.append(f"Expected exactly {expected_sheets}; found {workbook.sheetnames}")
    for sheet_name in expected_sheets:
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        rows, columns, headers = _sheet_shape_and_headers(worksheet)
        details[sheet_name] = {"rows": rows, "columns": columns}
        if rows != expected_rows:
            errors.append(f"{sheet_name} has {rows} rows; expected {expected_rows}")
        bad = _forbidden_headers(headers)
        if bad:
            errors.append(f"{sheet_name} contains excluded columns: {bad}")
    workbook.close()
    details["size_bytes"] = path.stat().st_size
    details["sha256"] = _sha256(path)
    return {"status": "PASS" if not errors else "FAIL", "errors": errors, "workbook": details}


def validate_step2(
    workbook_path: Path,
    manifest_path: Path,
    plot_root: Path,
    gene_rows: Sequence[dict[str, Any]],
    expected_width: int,
) -> dict[str, Any]:
    errors: list[str] = []
    expected_sheets = [SHEET_CELL_TYPES, SHEET_TISSUES, SHEET_GEP]
    workbook_details: dict[str, Any] = {"path": str(workbook_path), "exists": workbook_path.exists()}
    if not workbook_path.exists():
        errors.append(f"Missing final workbook: {workbook_path}")
    else:
        workbook = _workbook_api()(workbook_path, read_only=True, data_only=True)
        workbook_details["sheets"] = workbook.sheetnames
        if workbook.sheetnames != expected_sheets:
            errors.append(f"Expected final sheets {expected_sheets}; found {workbook.sheetnames}")
        for sheet_name in expected_sheets:
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            rows, columns, headers = _sheet_shape_and_headers(worksheet)
            workbook_details[sheet_name] = {"rows": rows, "columns": columns}
            if rows != len(gene_rows):
                errors.append(f"{sheet_name} has {rows} rows; expected {len(gene_rows)}")
            bad = _forbidden_headers(headers)
            if bad:
                errors.append(f"{sheet_name} contains excluded columns: {bad}")
        if SHEET_GEP in workbook.sheetnames:
            worksheet = workbook[SHEET_GEP]
            iterator = worksheet.iter_rows(values_only=True)
            try:
                headers = list(next(iterator))
            except StopIteration:
                headers = []
            expected_headers = [
                "Gene",
                "Ensembl",
                "GEP",
                "GEP_Zscore",
                "GEP_Loading",
                "GEP_Zscore_Rank",
            ]
            if headers != expected_headers:
                errors.append(f"Unexpected GEP_Info headers: {headers}")
            for index, values in enumerate(iterator):
                if index >= len(gene_rows):
                    break
                # Plain gene lists (no GEP columns) yield short rows; pad before comparing.
                values = tuple(values) + (None,) * max(0, 6 - len(values))
                if values[5] != gene_rows[index].get("GEP_Zscore_Rank"):
                    errors.append(
                        f"GEP rank mismatch at data row {index + 2}: "
                        f"{values[5]} != {gene_rows[index].get('GEP_Zscore_Rank')}"
                    )
                    break
        workbook.close()
        workbook_details["size_bytes"] = workbook_path.stat().st_size
        workbook_details["sha256"] = _sha256(workbook_path)

    plot_details: dict[str, Any] = {"manifest": str(manifest_path), "root": str(plot_root)}
    if not manifest_path.exists():
        errors.append(f"Missing plot manifest: {manifest_path}")
    else:
        with manifest_path.open("r", encoding="utf-8-sig", newline="") as handle:
            manifest_rows = list(csv.DictReader(handle))
        plot_details["manifest_rows"] = len(manifest_rows)
        if len(manifest_rows) != len(gene_rows):
            errors.append(f"Manifest has {len(manifest_rows)} rows; expected {len(gene_rows)}")
        try:
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError("Plot validation requires Pillow; install requirements.txt") from exc
        missing: list[str] = []
        corrupt: list[str] = []
        wrong_width: list[str] = []
        checked = 0
        for row in manifest_rows:
            for column in ("RNA_every_cell_type_path", "RNA_every_tissue_path"):
                path = plot_root / str(row[column])
                if not path.is_file():
                    missing.append(str(path))
                    continue
                checked += 1
                try:
                    with Image.open(path) as image:
                        if image.width != expected_width:
                            wrong_width.append(f"{path}: {image.width}")
                        image.verify()
                except Exception as exc:  # noqa: BLE001
                    corrupt.append(f"{path}: {exc}")
        expected_pngs = len(gene_rows) * 2
        plot_details.update(
            {
                "expected_pngs": expected_pngs,
                "checked_pngs": checked,
                "missing_count": len(missing),
                "corrupt_count": len(corrupt),
                "wrong_width_count": len(wrong_width),
                "examples": {
                    "missing": missing[:5],
                    "corrupt": corrupt[:5],
                    "wrong_width": wrong_width[:5],
                },
            }
        )
        if checked != expected_pngs:
            errors.append(f"Validated {checked} PNGs; expected {expected_pngs}")
        if missing:
            errors.append(f"Missing PNG files: {len(missing)}")
        if corrupt:
            errors.append(f"Corrupt PNG files: {len(corrupt)}")
        if wrong_width:
            errors.append(f"PNG files with unexpected width: {len(wrong_width)}")

    return {
        "status": "PASS" if not errors else "FAIL",
        "input_rows": len(gene_rows),
        "errors": errors,
        "workbook": workbook_details,
        "plots": plot_details,
    }
