from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence


SC_PREFIX = "Single Cell Type RNA - "
SC_SUFFIX = " [nCPM]"
TISSUE_PREFIX = "Tissue RNA - "
TISSUE_SUFFIX = " [nTPM]"

DEFAULT_CONFIG: dict[str, Any] = {
    "megakaryocytic_cell_types": [
        "Megakaryocyte-Erythroid progenitors",
        "Megakaryocyte progenitors",
        "Megakaryocytes",
        "Platelets",
    ],
    "hematopoietic_cell_types": [
        "Hematopoietic stem cells",
        "Erythrocyte progenitors",
        "Monocyte progenitors",
        "Neutrophil progenitors",
        "B-cells",
        "Erythrocytes",
        "Hofbauer cells",
        "Innate lymphoid cells",
        "Kupffer cells",
        "Macrophages",
        "Mast cells",
        "Microglia",
        "NK-cells",
        "Neutrophils",
        "Plasma cells",
        "T-cells",
        "Thymocytes",
        "cDC",
        "monocytes",
        "pDCs",
    ],
    "bone_marrow_non_hematopoietic_cell_types": [
        "Adipocytes",
        "Fibroblasts",
        "Pericytes",
        "Vascular endothelial cells",
        "Vascular smooth muscle cells",
    ],
    "hpa_base_url": "https://www.proteinatlas.org",
    "hpa_batch_size": 40,
    "hpa_request_retries": 4,
    "hpa_request_timeout_seconds": 120,
    "hpa_parallel_requests": 4,
    "plot_width_pixels": 1950,
    "plot_dpi": 150,
}


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def load_config(path: Path | None) -> dict[str, Any]:
    if path is None or not path.exists():
        return deepcopy(DEFAULT_CONFIG)
    with path.open("r", encoding="utf-8") as handle:
        return deep_merge(DEFAULT_CONFIG, json.load(handle))


def as_number(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    return parsed if math.isfinite(parsed) else None


def unique_preserve(values: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def _normalise_column_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _pick(row: dict[str, Any], names: Sequence[str]) -> Any:
    normalised = {_normalise_column_name(key): value for key, value in row.items()}
    for name in names:
        key = _normalise_column_name(name)
        if key in normalised and normalised[key] not in (None, ""):
            return normalised[key]
    return None


def _read_source_table(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        delimiter = "\t" if suffix in {".tsv", ".txt"} else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle, delimiter=delimiter)]
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel input requires openpyxl; install requirements.txt") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        iterator = worksheet.iter_rows(values_only=True)
        try:
            headers = [str(value or "") for value in next(iterator)]
        except StopIteration:
            workbook.close()
            return []
        rows = [dict(zip(headers, values)) for values in iterator]
        workbook.close()
        return rows
    raise ValueError(f"Unsupported gene-list format: {path.suffix}")


def read_gene_list(
    path: Path,
    requested_genes: set[str] | None = None,
    limit: int | None = None,
) -> list[dict[str, Any]]:
    source_rows = _read_source_table(path)
    rows: list[dict[str, Any]] = []
    for source_index, source in enumerate(source_rows, start=2):
        symbol = str(_pick(source, ["symbol", "gene symbol", "query_gene", "gene"]) or "").strip()
        if not symbol:
            continue
        if requested_genes is not None and symbol.upper() not in requested_genes:
            continue
        gep = str(_pick(source, ["top_GEP", "GEP"]) or "").strip() or None
        loading = as_number(_pick(source, ["top_GEP_loading", "loading", "loading_max"]))
        rows.append(
            {
                "Gene": symbol,
                "GEP": gep,
                "GEP_Zscore": as_number(_pick(source, ["top_GEP_zscore", "GEP_Zscore", "Zscore"])),
                "GEP_Loading": loading,
                "Source_Row": source_index,
            }
        )
        if limit is not None and len(rows) >= limit:
            break
    if not rows:
        raise ValueError(f"No usable gene symbols were found in {path}")
    return add_gep_zscore_ranks(rows)


def add_gep_zscore_ranks(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[tuple[int, float]]] = {}
    for index, row in enumerate(rows):
        gep = str(row.get("GEP") or "")
        zscore = as_number(row.get("GEP_Zscore"))
        if gep and zscore is not None:
            grouped.setdefault(gep, []).append((index, zscore))
    for entries in grouped.values():
        ordered = sorted(entries, key=lambda item: (-item[1], item[0]))
        previous: float | None = None
        current_rank = 0
        for position, (row_index, zscore) in enumerate(ordered, start=1):
            if previous is None or not math.isclose(zscore, previous, rel_tol=0.0, abs_tol=1e-12):
                current_rank = position
                previous = zscore
            rows[row_index]["GEP_Zscore_Rank"] = current_rank
    for row in rows:
        row.setdefault("GEP_Zscore_Rank", None)
    return rows


def _requests_module():
    try:
        import requests
    except ImportError as exc:
        raise RuntimeError("HPA download requires requests; install requirements.txt") from exc
    return requests


def discover_hpa_columns(base_url: str, timeout: int) -> tuple[list[str], list[str], str]:
    requests = _requests_module()
    url = f"{base_url.rstrip('/')}/search?search=gene_name%3ATP53"
    response = requests.get(
        url,
        timeout=timeout,
        headers={"User-Agent": "HPA-RNA-cluster-pipeline/2.0"},
    )
    response.raise_for_status()
    abbreviations = unique_preserve(re.findall(r'data-abbr="([^"]+)"', response.text))
    single_cell = [value for value in abbreviations if value.startswith("sc_RNA_") and value != "sc_RNA__tau"]
    tissues = [value for value in abbreviations if value.startswith("t_RNA_") and value != "t_RNA__tau"]
    if len(single_cell) < 100 or len(tissues) < 40:
        raise RuntimeError(
            "HPA export-column discovery returned an unexpected result: "
            f"single_cell={len(single_cell)}, tissues={len(tissues)}"
        )
    return single_cell, tissues, hashlib.sha256(response.content).hexdigest()


def _normalise_api_records(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return [record for record in payload if isinstance(record, dict)]
    if isinstance(payload, dict):
        for key in ("records", "data", "results"):
            value = payload.get(key)
            if isinstance(value, list):
                return [record for record in value if isinstance(record, dict)]
    raise RuntimeError("HPA API returned an unexpected JSON structure")


def _atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
    temporary.replace(path)


def _load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def _batch_cache_path(
    cache_dir: Path,
    dataset: str,
    batch_index: int,
    symbols: Sequence[str],
    columns: Sequence[str],
) -> Path:
    digest = hashlib.sha256(
        json.dumps([dataset, list(symbols), list(columns)], ensure_ascii=False).encode("utf-8")
    ).hexdigest()[:12]
    return cache_dir / dataset / f"batch_{batch_index:04d}_{digest}.json"


def _fetch_hpa_batch(
    base_url: str,
    symbols: list[str],
    columns: list[str],
    dataset: str,
    batch_index: int,
    retries: int,
    timeout: int,
    cache_path: Path,
    refresh: bool,
) -> tuple[str, int, list[dict[str, Any]]]:
    if cache_path.exists() and not refresh:
        return dataset, batch_index, _normalise_api_records(_load_json(cache_path))
    requests = _requests_module()
    endpoint = f"{base_url.rstrip('/')}/api/search_download.php"
    params = {
        "search": "gene_name:" + ",".join(symbols),
        "format": "json",
        "columns": ",".join(columns),
        "compress": "no",
    }
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            response = requests.get(
                endpoint,
                params=params,
                timeout=timeout,
                headers={"User-Agent": "HPA-RNA-cluster-pipeline/2.0"},
            )
            response.raise_for_status()
            records = _normalise_api_records(response.json())
            _atomic_json(cache_path, records)
            return dataset, batch_index, records
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            if attempt < retries:
                time.sleep(min(30, attempt * 3))
    raise RuntimeError(f"HPA {dataset} batch {batch_index} failed: {last_error}")


def _ordered_names(all_keys: set[str], prefix: str, suffix: str = "") -> list[str]:
    names: list[str] = []
    for key in all_keys:
        if key.startswith(prefix) and (not suffix or key.endswith(suffix)):
            end = -len(suffix) if suffix else None
            names.append(key[len(prefix):end])
    return sorted(set(names), key=str.casefold)


def _order_cell_types(names: Sequence[str], config: dict[str, Any]) -> list[str]:
    featured = unique_preserve(
        list(config["megakaryocytic_cell_types"])
        + list(config["hematopoietic_cell_types"])
        + list(config["bone_marrow_non_hematopoietic_cell_types"])
    )
    available = set(names)
    featured_set = set(featured)
    return [name for name in featured if name in available] + [
        name for name in names if name not in featured_set
    ]


def _order_tissues(names: Sequence[str]) -> list[str]:
    return (["bone marrow"] if "bone marrow" in names else []) + [
        name for name in names if name != "bone marrow"
    ]


def _matrix(
    gene_rows: Sequence[dict[str, Any]],
    records: dict[str, dict[str, Any]],
    names: Sequence[str],
    prefix: str,
    suffix: str,
) -> dict[str, Any]:
    rows: list[list[Any]] = []
    for item in gene_rows:
        gene = str(item["Gene"])
        record = records.get(gene.upper(), {"Gene": gene})
        rows.append(
            [gene, record.get("Ensembl")]
            + [as_number(record.get(f"{prefix}{name}{suffix}")) for name in names]
        )
    return {"headers": ["Gene", "Ensembl"] + list(names), "rows": rows}


def fetch_hpa_rna(
    gene_rows: list[dict[str, Any]],
    config: dict[str, Any],
    cache_dir: Path,
    refresh: bool = False,
) -> dict[str, Any]:
    symbols = unique_preserve(str(row["Gene"]) for row in gene_rows)
    base_url = str(config["hpa_base_url"])
    timeout = int(config["hpa_request_timeout_seconds"])
    sc_columns, tissue_columns, page_hash = discover_hpa_columns(base_url, timeout)
    datasets = {
        "single_cell": unique_preserve(["g", "eg"] + sc_columns),
        "tissues": unique_preserve(["g", "eg"] + tissue_columns),
    }
    batch_size = int(config["hpa_batch_size"])
    batches = [symbols[index:index + batch_size] for index in range(0, len(symbols), batch_size)]
    jobs: list[tuple[Any, ...]] = []
    for dataset, columns in datasets.items():
        for batch_index, batch in enumerate(batches, start=1):
            jobs.append(
                (
                    base_url,
                    batch,
                    columns,
                    dataset,
                    batch_index,
                    int(config["hpa_request_retries"]),
                    timeout,
                    _batch_cache_path(cache_dir, dataset, batch_index, batch, columns),
                    refresh,
                )
            )
    records_by_dataset: dict[str, dict[str, dict[str, Any]]] = {
        "single_cell": {},
        "tissues": {},
    }
    print(f"HPA RNA download: {len(symbols)} unique genes, {len(jobs)} batch requests", flush=True)
    with ThreadPoolExecutor(max_workers=int(config["hpa_parallel_requests"])) as pool:
        futures = [pool.submit(_fetch_hpa_batch, *job) for job in jobs]
        for completed, future in enumerate(as_completed(futures), start=1):
            dataset, _, records = future.result()
            for record in records:
                gene = str(record.get("Gene") or "").upper()
                if gene:
                    records_by_dataset[dataset][gene] = record
            if completed == 1 or completed % 10 == 0 or completed == len(jobs):
                print(f"HPA batches completed: {completed}/{len(jobs)}", flush=True)

    sc_keys = {key for record in records_by_dataset["single_cell"].values() for key in record}
    tissue_keys = {key for record in records_by_dataset["tissues"].values() for key in record}
    sc_names = _order_cell_types(_ordered_names(sc_keys, SC_PREFIX, SC_SUFFIX), config)
    tissue_names = _order_tissues(_ordered_names(tissue_keys, TISSUE_PREFIX, TISSUE_SUFFIX))
    if not sc_names or not tissue_names:
        raise RuntimeError("HPA returned no usable single-cell or tissue RNA columns")

    missing = [
        gene for gene in symbols
        if gene.upper() not in records_by_dataset["single_cell"]
        and gene.upper() not in records_by_dataset["tissues"]
    ]
    return {
        "metadata": {
            "source": "Human Protein Atlas public API",
            "source_url": f"{base_url.rstrip('/')}/api/search_download.php",
            "retrieved_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "gene_rows": len(gene_rows),
            "unique_genes": len(symbols),
            "missing_genes": missing,
            "single_cell_types": len(sc_names),
            "tissues": len(tissue_names),
            "search_page_sha256": page_hash,
        },
        "single_cell": _matrix(
            gene_rows,
            records_by_dataset["single_cell"],
            sc_names,
            SC_PREFIX,
            SC_SUFFIX,
        ),
        "tissues": _matrix(
            gene_rows,
            records_by_dataset["tissues"],
            tissue_names,
            TISSUE_PREFIX,
            TISSUE_SUFFIX,
        ),
    }


def _default_config_path() -> Path:
    return Path(__file__).resolve().parents[1] / "config.json"


def _requested_set(value: str | None) -> set[str] | None:
    if not value:
        return None
    return {item.strip().upper() for item in value.split(",") if item.strip()}


def step1_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Step 1: download HPA single-cell and tissue RNA expression into a two-sheet Excel file."
    )
    parser.add_argument("--gene-list", type=Path, required=True, help="CSV/TSV/XLSX gene list")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--output-excel", type=Path, help="Default: OUTPUT_DIR/HPA_RNA_expression.xlsx")
    parser.add_argument("--config", type=Path, default=_default_config_path())
    parser.add_argument("--refresh", action="store_true", help="Re-download batches even if cache files exist")
    parser.add_argument("--genes", help="Optional comma-separated subset for testing")
    parser.add_argument("--limit", type=int, help="Optional first-N rows for testing")
    return parser


def step1_main(argv: Sequence[str] | None = None) -> int:
    args = step1_parser().parse_args(argv)
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    config = load_config(args.config)
    rows = read_gene_list(args.gene_list.resolve(), _requested_set(args.genes), args.limit)
    payload = fetch_hpa_rna(rows, config, output_dir / "cache", refresh=args.refresh)
    output_excel = (args.output_excel or output_dir / "HPA_RNA_expression.xlsx").resolve()
    from .excel_report import write_expression_workbook
    from .validate import validate_step1

    write_expression_workbook(payload, output_excel)
    validation = validate_step1(output_excel, expected_rows=len(rows))
    validation_path = output_dir / "validation_step1.json"
    validation_path.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    if validation["status"] != "PASS":
        raise RuntimeError(f"Step 1 validation failed: {validation['errors']}")
    print(json.dumps(payload["metadata"], ensure_ascii=False, indent=2), flush=True)
    print(f"STEP1_DONE {output_excel}", flush=True)
    return 0


def step2_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Step 2: draw two RNA plots per gene and merge GEP fields into a final Excel file."
    )
    parser.add_argument("--gene-list", type=Path, required=True, help="The same or updated gene-list format")
    parser.add_argument("--hpa-excel", type=Path, required=True, help="Excel file produced by step 1")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--output-excel", type=Path, help="Default: OUTPUT_DIR/HPA_RNA_with_GEP.xlsx")
    parser.add_argument("--config", type=Path, default=_default_config_path())
    parser.add_argument("--workers", type=int, default=1, help="Parallel plot processes")
    parser.add_argument("--genes", help="Optional comma-separated subset for testing")
    parser.add_argument("--limit", type=int, help="Optional first-N rows for testing")
    return parser


def step2_main(argv: Sequence[str] | None = None) -> int:
    args = step2_parser().parse_args(argv)
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    config = load_config(args.config)
    rows = read_gene_list(args.gene_list.resolve(), _requested_set(args.genes), args.limit)
    from .excel_report import read_expression_workbook, write_final_workbook
    from .plots import generate_plots
    from .validate import validate_step2

    expression = read_expression_workbook(args.hpa_excel.resolve())
    manifest = generate_plots(
        expression,
        rows,
        output_dir / "gene_plots",
        config,
        workers=max(1, args.workers),
    )
    output_excel = (args.output_excel or output_dir / "HPA_RNA_with_GEP.xlsx").resolve()
    write_final_workbook(expression, rows, output_excel)
    validation = validate_step2(
        output_excel,
        manifest,
        output_dir / "gene_plots",
        rows,
        int(config.get("plot_width_pixels", 1950)),
    )
    validation_path = output_dir / "validation_step2.json"
    validation_path.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    if validation["status"] != "PASS":
        raise RuntimeError(f"Step 2 validation failed: {validation['errors']}")
    print(json.dumps(validation, ensure_ascii=False, indent=2), flush=True)
    print(f"STEP2_DONE {output_excel}", flush=True)
    return 0


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Two-step HPA RNA workflow")
    parser.add_argument("command", choices=["step1", "step2"])
    known, remaining = parser.parse_known_args(argv)
    if known.command == "step1":
        return step1_main(remaining)
    return step2_main(remaining)
