#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Screen a gene list for cell-type-specific expression using the HPA workbook from step 1.

Rule (learned from real usage): the single-cell-type peak is the criterion.
For each gene, find the highest-expressing HPA cell type; the gene is selected when
that peak type belongs to one of the target groups defined in the config
(`specificity_groups`, ordered by priority — the first matching group is the driver).
Tissue values are annotation only, never a filter: a gene can be cell-type specific
without being tissue specific.

Borderline: genes whose best target-group cell type reaches >= borderline_fraction
(default 0.85) of the overall peak are reported separately for manual review.

Usage:
    python screen_specificity.py --hpa-excel HPA_RNA_expression.xlsx \
        --config config.json --out-dir results/ [--membership groups.tsv]

membership file (optional): two columns `group<TAB>gene`, adds a membership column
to the output (e.g. GEP assignment). Duplicate genes keep all groups, `;`-joined.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def read_sheet(ws) -> dict[str, dict[str, float | None]]:
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    out: dict[str, dict[str, float | None]] = {}
    for row in it:
        if not row or row[0] is None or row[0] in out:
            continue
        out[str(row[0]).strip()] = {
            header[j]: (float(row[j]) if j < len(row) and row[j] not in (None, "", "N/A") else None)
            for j in range(2, len(header))
        }
    return out


def load_membership(path: Path) -> dict[str, list[str]]:
    m: dict[str, list[str]] = {}
    for ln in path.read_text(encoding="utf-8").splitlines()[1:]:
        parts = ln.split("\t")
        if len(parts) >= 2 and parts[1].strip():
            grp, gene = parts[0].strip(), parts[1].strip()
            m.setdefault(gene, [])
            if grp not in m[gene]:
                m[gene].append(grp)
    return m


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Screen genes for cell-type-specific expression (HPA).")
    ap.add_argument("--hpa-excel", type=Path, required=True, help="Workbook from step1_fetch_hpa.py")
    ap.add_argument("--config", type=Path, required=True,
                    help="JSON with `specificity_groups` (ordered dict of group -> cell types); "
                         "optional `reference_tissue`, `borderline_fraction`")
    ap.add_argument("--out-dir", type=Path, required=True)
    ap.add_argument("--membership", type=Path, help="Optional two-column group<TAB>gene TSV")
    args = ap.parse_args(argv)

    cfg = json.loads(args.config.read_text(encoding="utf-8"))
    groups: dict[str, set[str]] = {k: set(v) for k, v in cfg["specificity_groups"].items()}
    if not groups:
        raise SystemExit("config must define non-empty `specificity_groups`")
    ref_tissue = cfg.get("reference_tissue")
    borderline_fraction = float(cfg.get("borderline_fraction", 0.85))
    union = set().union(*groups.values())

    args.out_dir.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(args.hpa_excel, read_only=True)
    cells = read_sheet(wb["RNA_Cell_Types"])
    tissue = read_sheet(wb["RNA_Tissues"]) if "RNA_Tissues" in wb.sheetnames else {}
    wb.close()
    membership = load_membership(args.membership) if args.membership else {}

    selected, borderline = [], []
    for gene, cv in cells.items():
        if any(v is None for v in cv.values()):
            continue  # unmatched gene: row left blank upstream
        grp_vals = [(c, v) for c, v in cv.items() if c in union and v is not None]
        if not grp_vals:
            continue
        top_cell = max(cv, key=lambda c: cv[c] if cv[c] is not None else -1)
        grp_max = max(v for _, v in grp_vals)
        oth_vals = [v for c, v in cv.items() if c not in union and v is not None]
        oth_max = max(oth_vals) if oth_vals else 0.0
        driver = next((name for name, members in groups.items() if top_cell in members), None)

        tv = tissue.get(gene, {})
        bm = tv.get(ref_tissue) if ref_tissue else None
        bm_rank = None
        if bm is not None:
            bm_rank = 1 + sum(1 for c, v in tv.items()
                              if c != ref_tissue and v is not None and v > bm)

        rec = {
            "gene": gene,
            "membership": ";".join(membership.get(gene, [])),
            "driver": driver or "",
            "top_cell_type": top_cell,
            "top_cell_value": cv[top_cell],
            "group_max": grp_max,
            "other_max": oth_max,
            "specificity_ratio": (grp_max / oth_max) if oth_max else None,
            "ref_tissue_value": bm,
            "ref_tissue_rank": bm_rank,
        }
        if driver and grp_max > 0 and grp_max >= oth_max:
            selected.append(rec)
        elif grp_max > 0 and oth_max and grp_max >= borderline_fraction * oth_max:
            borderline.append(rec)

    order = {name: i for i, name in enumerate(groups)}
    selected.sort(key=lambda r: (order[r["driver"]], -(r["group_max"]), r["gene"]))
    borderline.sort(key=lambda r: -(r["group_max"]))

    tsv = args.out_dir / "selected_genes.tsv"
    with tsv.open("w", encoding="utf-8") as f:
        f.write("gene\tmembership\tdriver\ttop_cell_type\ttop_cell_value\tgroup_max\tother_max"
                "\tspecificity_ratio\tref_tissue_value\tref_tissue_rank\n")
        for r in selected:
            f.write("\t".join(str(r[k]) if r[k] is not None else "" for k in
                              ("gene", "membership", "driver", "top_cell_type", "top_cell_value",
                               "group_max", "other_max", "specificity_ratio",
                               "ref_tissue_value", "ref_tissue_rank")) + "\n")

    summary = {
        "selected": len(selected),
        "by_driver": {name: sum(1 for r in selected if r["driver"] == name) for name in groups},
        "borderline": len(borderline),
        "borderline_genes": [r["gene"] for r in borderline],
        "config": {"specificity_groups": {k: sorted(v) for k, v in groups.items()},
                   "reference_tissue": ref_tissue,
                   "borderline_fraction": borderline_fraction},
    }
    (args.out_dir / "screening_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"selected: {len(selected)} {summary['by_driver']}; borderline: {len(borderline)}")
    print(f"SCREEN_DONE {tsv}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
