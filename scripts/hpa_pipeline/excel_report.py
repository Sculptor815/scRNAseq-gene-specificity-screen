from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, Sequence


SHEET_CELL_TYPES = "RNA_Cell_Types"
SHEET_TISSUES = "RNA_Tissues"
SHEET_GEP = "GEP_Info"

HEADER_FILL = "0F6B6D"
HEADER_FONT = "FFFFFF"
TEXT_COLOR = "1F2937"
LINE_COLOR = "C9D3DC"


def _imports():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel output requires openpyxl; install requirements.txt") from exc
    return {
        "Workbook": Workbook,
        "load_workbook": load_workbook,
        "WriteOnlyCell": WriteOnlyCell,
        "Alignment": Alignment,
        "Border": Border,
        "Font": Font,
        "PatternFill": PatternFill,
        "Side": Side,
        "get_column_letter": get_column_letter,
    }


def _header_cells(worksheet, headers: Sequence[Any], api: dict[str, Any]) -> list[Any]:
    fill = api["PatternFill"]("solid", fgColor=HEADER_FILL)
    font = api["Font"](bold=True, color=HEADER_FONT)
    alignment = api["Alignment"](horizontal="center", vertical="center", wrap_text=True)
    border = api["Border"](bottom=api["Side"](style="thin", color=LINE_COLOR))
    cells = []
    for value in headers:
        cell = api["WriteOnlyCell"](worksheet, value=value)
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment
        cell.border = border
        cells.append(cell)
    return cells


def _write_matrix_sheet(
    workbook,
    name: str,
    headers: Sequence[Any],
    rows: Iterable[Sequence[Any]],
    frozen_column: int,
    api: dict[str, Any],
    number_formats: dict[int, str] | None = None,
) -> int:
    worksheet = workbook.create_sheet(name)
    worksheet.sheet_view.showGridLines = True
    worksheet.freeze_panes = f"{api['get_column_letter'](frozen_column + 1)}2"
    worksheet.sheet_view.zoomScale = 80
    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 24
    for index in range(3, len(headers) + 1):
        worksheet.column_dimensions[api["get_column_letter"](index)].width = 15
    header_widths = {
        "GEP": 15,
        "GEP_Zscore": 19,
        "GEP_Loading": 19,
        "GEP_Zscore_Rank": 20,
    }
    for index, header in enumerate(headers, start=1):
        if str(header) in header_widths:
            worksheet.column_dimensions[api["get_column_letter"](index)].width = header_widths[str(header)]
    worksheet.append(_header_cells(worksheet, headers, api))
    row_count = 1
    formats = number_formats or {}
    for row in rows:
        values = list(row)
        if formats:
            output_row: list[Any] = []
            for column_index, value in enumerate(values, start=1):
                number_format = formats.get(column_index)
                if number_format and isinstance(value, (int, float)):
                    cell = api["WriteOnlyCell"](worksheet, value=value)
                    cell.number_format = number_format
                    output_row.append(cell)
                else:
                    output_row.append(value)
            worksheet.append(output_row)
        else:
            worksheet.append(values)
        row_count += 1
    worksheet.auto_filter.ref = f"A1:{api['get_column_letter'](len(headers))}{row_count}"
    return row_count - 1


def _atomic_save(workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(output_path.stem + ".tmp" + output_path.suffix)
    workbook.save(temporary)
    temporary.replace(output_path)


def write_expression_workbook(payload: dict[str, Any], output_path: Path) -> None:
    api = _imports()
    workbook = api["Workbook"](write_only=True)
    _write_matrix_sheet(
        workbook,
        SHEET_CELL_TYPES,
        payload["single_cell"]["headers"],
        payload["single_cell"]["rows"],
        frozen_column=2,
        api=api,
    )
    _write_matrix_sheet(
        workbook,
        SHEET_TISSUES,
        payload["tissues"]["headers"],
        payload["tissues"]["rows"],
        frozen_column=2,
        api=api,
    )
    workbook.properties.title = "Human Protein Atlas RNA expression"
    workbook.properties.subject = "Single-cell nCPM and tissue nTPM"
    workbook.properties.description = (
        "Raw HPA RNA expression matrices only. No IHC, candidate score, or derived specificity metric."
    )
    _atomic_save(workbook, output_path)


def _read_sheet(workbook, sheet_name: str) -> dict[str, Any]:
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = list(next(iterator))
    except StopIteration:
        return {"headers": [], "rows": []}
    rows: list[list[Any]] = []
    for row in iterator:
        values = list(row)
        if len(values) < len(headers):
            values.extend([None] * (len(headers) - len(values)))
        rows.append(values)
    return {"headers": headers, "rows": rows}


def read_expression_workbook(path: Path) -> dict[str, Any]:
    api = _imports()
    workbook = api["load_workbook"](path, read_only=True, data_only=True)
    expected = [SHEET_CELL_TYPES, SHEET_TISSUES]
    if workbook.sheetnames != expected:
        workbook.close()
        raise ValueError(f"Expected step-1 sheets {expected}, found {workbook.sheetnames}")
    payload = {
        "single_cell": _read_sheet(workbook, SHEET_CELL_TYPES),
        "tissues": _read_sheet(workbook, SHEET_TISSUES),
    }
    workbook.close()
    return payload


def _rows_by_gene(matrix: dict[str, Any]) -> dict[str, list[Any]]:
    return {
        str(row[0]).upper(): list(row)
        for row in matrix["rows"]
        if row and row[0] not in (None, "")
    }


def _metadata(row: dict[str, Any], ensembl: Any) -> list[Any]:
    return [
        row.get("Gene"),
        ensembl,
        row.get("GEP"),
        row.get("GEP_Zscore"),
        row.get("GEP_Loading"),
        row.get("GEP_Zscore_Rank"),
    ]


def _final_matrix_rows(
    gene_rows: Sequence[dict[str, Any]],
    matrix: dict[str, Any],
) -> Iterable[list[Any]]:
    by_gene = _rows_by_gene(matrix)
    empty_values = [None] * max(0, len(matrix["headers"]) - 2)
    for row in gene_rows:
        expression = by_gene.get(str(row["Gene"]).upper())
        ensembl = expression[1] if expression else None
        values = expression[2:] if expression else empty_values
        yield _metadata(row, ensembl) + list(values)


def write_final_workbook(
    expression: dict[str, Any],
    gene_rows: Sequence[dict[str, Any]],
    output_path: Path,
) -> None:
    api = _imports()
    workbook = api["Workbook"](write_only=True)
    metadata_headers = [
        "Gene",
        "Ensembl",
        "GEP",
        "GEP_Zscore",
        "GEP_Loading",
        "GEP_Zscore_Rank",
    ]
    _write_matrix_sheet(
        workbook,
        SHEET_CELL_TYPES,
        metadata_headers + list(expression["single_cell"]["headers"][2:]),
        _final_matrix_rows(gene_rows, expression["single_cell"]),
        frozen_column=6,
        api=api,
        number_formats={4: "0.000", 5: "0.000", 6: "0"},
    )
    _write_matrix_sheet(
        workbook,
        SHEET_TISSUES,
        metadata_headers + list(expression["tissues"]["headers"][2:]),
        _final_matrix_rows(gene_rows, expression["tissues"]),
        frozen_column=6,
        api=api,
        number_formats={4: "0.000", 5: "0.000", 6: "0"},
    )

    cell_map = _rows_by_gene(expression["single_cell"])
    tissue_map = _rows_by_gene(expression["tissues"])
    gep_rows = []
    for row in gene_rows:
        gene = str(row["Gene"]).upper()
        cell_record = cell_map.get(gene)
        tissue_record = tissue_map.get(gene)
        ensembl = (cell_record[1] if cell_record else None) or (tissue_record[1] if tissue_record else None)
        gep_rows.append(_metadata(row, ensembl))
    _write_matrix_sheet(
        workbook,
        SHEET_GEP,
        metadata_headers,
        gep_rows,
        frozen_column=2,
        api=api,
        number_formats={4: "0.000", 5: "0.000", 6: "0"},
    )
    workbook.properties.title = "HPA RNA expression with GEP annotations"
    workbook.properties.description = (
        "HPA RNA matrices plus input GEP, z-score, loading, and within-GEP z-score rank. "
        "No additional score or derived expression metric."
    )
    _atomic_save(workbook, output_path)
